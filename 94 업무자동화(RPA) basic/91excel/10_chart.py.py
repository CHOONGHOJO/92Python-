from openpyxl import load_workbook
wb = load_workbook("sample5.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, Reference, LineChart
# B2:C11 까지의 데이타를 쳐트로 작성
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart() # 차트 종류 설정 (Bar, Line, Pie, ...)
# bar_chart.add_data(bar_value) # 차트 데이타 추가

# B1:C11 까지의 데이타를 쳐트로 작성
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart() # 차트 종류 설정 (Bar, Line, Pie, ...)
line_chart.add_data(line_value, titles_from_data=True) # 계열 > 영어, 수학 (제목에서 가져옴)
line_chart.title = "성적표" # 제목
line_chart.style = 10 # 미리 정의된 스타일 적용, 사용자 개별 지정 가능
line_chart.y_axis.title = "점수" # y축의 제목
line_chart.x_axis.title = "번호" # x축의 제목
ws.add_chart(line_chart, "E1") # 차트 넣을 위치 정의



wb.save("sample10_chart.xlsx")